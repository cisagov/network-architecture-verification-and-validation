"""CLI Commands."""
import os
import webbrowser
import json
import shutil
import sys

# Third-Party Libraries
import click
import openpyxl

# cisagov Libraries
from navv.gui.app import app
from navv.bll import (
    get_snmp_df,
    get_zeek_df,
    get_mac_df,
    get_http_df,
    get_ssl_df,
    get_generic_df,
)
from navv.message_handler import success_msg, warning_msg, error_msg
from navv.spreadsheet_tools import (
    auto_adjust_width,
    create_target_workbook,
    write_inventory_sheet,
    read_existing_notes,
    create_analysis_array,
    get_inventory_data,
    get_package_data,
    get_segments_data,
    get_workbook,
    perform_analysis,
    write_conn_states_sheet,
    write_externals_sheet,
    write_snmp_sheet,
    write_stats_sheet,
    write_unknown_internals_sheet,
    write_internal_hosts_sheet,
    write_ipv6_hosts_sheet,
    write_purdue_violations_sheet,
    write_legend_sheet,
    write_data_layer_sheet,
    generate_sankey_html,
    auto_discover_segments,
    write_segments_sheet,
    color_inventory_sheet,
    write_external_inbound_sheet,
    write_internal_outbound_sheet,
    write_zeek_log_sheets,
)
from navv.zeek import (
    get_conn_data,
    get_dns_data,
    get_dhcp_data,
    get_snmp_data,
    get_http_data,
    get_ssl_data,
    get_log_data,
    run_zeek,
    perform_zeekcut,
)
from navv.utilities import pushd, get_mac_vendor
from navv.geolocation import Geolocator
from navv.elevadr import generate_elevadr_report


@click.command("generate")
@click.option(
    "-o",
    "--output-dir",
    required=False,
    help="Directory to place resultant analysis files in. Defaults to current working directory.",
    type=str,
    default=os.getcwd(),
)
@click.option(
    "-p",
    "--pcap",
    required=False,
    help="Path to pcap file. NAVV requires zeek logs or pcap. If used, zeek will run on pcap to create new logs.",
    type=str,
)
@click.option(
    "-z",
    "--zeek-logs",
    required=False,
    help="Path to store or contain zeek log files. Defaults to current working directory.",
    type=str,
    default=os.getcwd(),
)
@click.option(
    "-g",
    "--geoip-db",
    required=False,
    help="Path to GeoLite2 database file or directory (MMDB format). If not specified, DB-IP Lite databases will be automatically downloaded and cached.",
    type=str,
)
@click.option(
    "-m",
    "--macro",
    is_flag=True,
    help="Use macro-enabled template MACRO_network_analysis.xlsm as base.",
)
@click.argument("customer_name")
def generate(customer_name, output_dir, pcap, zeek_logs, geoip_db, macro):
    """Generate excel sheet."""
    if not shutil.which("zeek"):
        msg = "Zeek is not installed or not found in PATH. Please install Zeek to use NAVV."
        if sys.platform == "win32":
            msg += " Note: On Windows, NAVV must be run from within the WSL terminal where Zeek is installed."
        error_msg(msg)
        sys.exit(1)
    if not shutil.which("zeek-cut"):
        msg = "zeek-cut is not installed or not found in PATH. Please install Zeek to use NAVV."
        if sys.platform == "win32":
            msg += " Note: On Windows, NAVV must be run from within the WSL terminal where Zeek is installed."
        error_msg(msg)
        sys.exit(1)

    with pushd(output_dir):
        pass
    ext = ".xlsm" if macro else ".xlsx"
    file_name = os.path.join(output_dir, customer_name + "_network_analysis" + ext)

    # Read from existing workbook if it exists
    segments = []
    inventory = {}
    existing_notes = {}
    inventory_tab_name = "Inventory Input"

    other_ext = ".xlsx" if macro else ".xlsm"
    file_name_to_read = None
    if os.path.isfile(file_name):
        file_name_to_read = file_name
    else:
        fallback_file_name = os.path.join(output_dir, customer_name + "_network_analysis" + other_ext)
        if os.path.isfile(fallback_file_name):
            file_name_to_read = fallback_file_name

    if file_name_to_read:
        existing_wb = openpyxl.load_workbook(file_name_to_read)
        segments = get_segments_data(existing_wb["Segments"])
        inventory_tab_name = "Inventory" if "Inventory" in existing_wb.sheetnames else "Inventory Input"
        inventory = get_inventory_data(existing_wb[inventory_tab_name])
        existing_notes = read_existing_notes(existing_wb)
        existing_wb.close()

    # Create target workbook
    wb = create_target_workbook(macro=macro, inventory_tab_name=inventory_tab_name)

    # Initialize geolocator for IP geolocation
    geolocator = Geolocator(db_path=geoip_db)

    services, conn_states = get_package_data()
    timer_data = dict()

    if pcap:
        try:
            run_zeek(os.path.abspath(pcap), zeek_logs, timer=timer_data)
        except Exception as e:
            error_msg(f"Zeek failed to run on PCAP. Please verify that the PCAP file is not corrupt and that Zeek is installed and working correctly: {e}")
            sys.exit(1)
    else:
        timer_data["run_zeek"] = "NOT RAN"

    # Get zeek data from conn.log, dns.log and snmp.log
    zeek_data = get_conn_data(zeek_logs)
    snmp_data = get_snmp_data(zeek_logs)
    dns_filtered = get_dns_data(customer_name, output_dir, zeek_logs)
    dhcp_data = get_dhcp_data(zeek_logs)
    http_data = get_http_data(zeek_logs)
    ssl_data = get_ssl_data(zeek_logs)
    
    # Merge dhcp hostnames into dns dictionary
    for ip, hostname in dhcp_data.items():
        if ip not in dns_filtered:
            dns_filtered[ip] = hostname

    # Get dns data for resolution
    json_path = os.path.join(output_dir, f"{customer_name}_dns_data.json")
    
    # Get external DNS cache (similar to dns_data pattern)
    ext_dns_path = os.path.join(output_dir, f"{customer_name}_ext_dns_cache.json")
    if os.path.exists(ext_dns_path):
        with open(ext_dns_path, "r", encoding="utf-8") as f:
            try:
                ext_dns_cache = json.load(f)
            except:
                ext_dns_cache = {}
    else:
        ext_dns_cache = {}

    # Get zeek dataframes
    zeek_df = get_zeek_df(zeek_data, dns_filtered)
    snmp_df = get_snmp_df(snmp_data)
    http_df = get_http_df(http_data)
    ssl_df = get_ssl_df(ssl_data)
    
    zeek_dfs = {
        "HTTP": http_df,
        "SSL": ssl_df
    }

    # Additional OT / Remote Access / DNS Logs
    extra_logs = {
        "DNS": ("dns", ["id.orig_h", "id.resp_h", "id.resp_p", "proto", "query", "qclass_name", "qtype_name", "rcode_name", "answers"]),
        "Modbus": ("modbus", ["id.orig_h", "id.resp_h", "id.resp_p", "func", "exception"]),
        "DNP3": ("dnp3", ["id.orig_h", "id.resp_h", "id.resp_p", "fc_request", "fc_reply", "iin"]),
        "BACnet": ("bacnet", ["id.orig_h", "id.resp_h", "id.resp_p", "pdu_service"]),
        "ENIP": ("enip", ["id.orig_h", "id.resp_h", "id.resp_p", "command"]),
        "SSH": ("ssh", ["id.orig_h", "id.resp_h", "id.resp_p", "auth_success", "client", "server", "cipher_alg"]),
        "RDP": ("rdp", ["id.orig_h", "id.resp_h", "id.resp_p", "cookie", "result", "security_protocol"]),
    }
    
    for sheet_name, (log_name, fields) in extra_logs.items():
        log_data = get_log_data(zeek_logs, log_name, fields)
        if log_data and len(log_data) > 0 and log_data[0] != "":
            zeek_dfs[sheet_name] = get_generic_df(log_data, fields)

    # Get mac dataframe
    mac_df = get_mac_df(zeek_df)

    import pandas as pd
    smac_df = zeek_df[['src_mac', 'src_ip']].rename(columns={'src_mac': 'mac', 'src_ip': 'ip'})
    dmac_df = zeek_df[['dst_mac', 'dst_ip']].rename(columns={'dst_mac': 'mac', 'dst_ip': 'ip'})
    all_macs = pd.concat([smac_df, dmac_df], ignore_index=True)
    all_macs = all_macs[(all_macs['mac'].notna()) & (all_macs['mac'] != '-')]
    
    all_macs['is_ipv4'] = all_macs['ip'].apply(lambda x: 1 if ':' not in str(x) else 0)
    
    ipv4_counts = all_macs[all_macs['is_ipv4'] == 1].groupby('mac')['ip'].nunique()
    ipv6_counts = all_macs[all_macs['is_ipv4'] == 0].groupby('mac')['ip'].nunique()
    ip_to_macs = all_macs.groupby('ip')['mac'].unique()
    
    MAC_VENDORS_JSON_FILE = os.path.abspath(__file__ + "/../" + "data/mac-vendors.json")
    with open(MAC_VENDORS_JSON_FILE, encoding="utf-8") as f:
        mac_vendors = json.load(f)
        
    ip_to_mac_label = {}
    for ip, macs in ip_to_macs.items():
        labels = []
        for m in macs:
            if ipv4_counts.get(m, 0) > 1 or ipv6_counts.get(m, 0) > 1:
                if "Network Device" not in labels:
                    labels.append("Network Device")
            else:
                vendor = get_mac_vendor(mac_vendors, m.strip())
                if vendor and vendor != "Unknown vendor":
                    labels.append(f"{m} ({vendor})")
                else:
                    labels.append(m)
        ip_to_mac_label[ip] = ", ".join(labels)

    # Turn zeekcut data into rows for spreadsheet
    rows = create_analysis_array(zeek_data, timer=timer_data)

    # Auto-discover and recolor segments
    segments = auto_discover_segments(zeek_df, segments)
    write_segments_sheet(segments, wb)
    write_inventory_sheet(inventory, wb, inventory_tab_name)
    color_inventory_sheet(wb, inventory_tab_name, segments)

    ext_IPs = set()
    unk_int_IPs = set()
    purdue_violations = []
    sankey_data = {}
    verified_sankey_data = {}
    macro_sankey_data = {}
    verified_macro_sankey_data = {}
    macro_link_colors = {}
    perform_analysis(
        wb,
        rows,
        services,
        conn_states,
        inventory,
        segments,
        dns_filtered,
        json_path,
        ext_IPs,
        unk_int_IPs,
        purdue_violations=purdue_violations,
        sankey_data=sankey_data,
        verified_sankey_data=verified_sankey_data,
        macro_sankey_data=macro_sankey_data,
        verified_macro_sankey_data=verified_macro_sankey_data,
        macro_link_colors=macro_link_colors,
        geolocator=geolocator,
        ext_dns_cache=ext_dns_cache,
        ip_to_mac_label=ip_to_mac_label,
        timer=timer_data,
        existing_notes=existing_notes,
    )

    click.echo("Writing externals, unknown internals, SNMP, internal hosts, and Purdue violations sheets...")
    write_externals_sheet(ext_IPs, wb, geolocator=geolocator, ext_dns_cache=ext_dns_cache)

    write_unknown_internals_sheet(unk_int_IPs, wb)

    write_snmp_sheet(snmp_df, wb)

    write_internal_hosts_sheet(mac_df, wb, inventory, segments)
    
    write_ipv6_hosts_sheet(mac_df, wb)
    
    write_purdue_violations_sheet(purdue_violations, wb)
    
    click.echo("Writing inbound, outbound, and specialized Zeek log sheets...")
    write_external_inbound_sheet(rows, wb)
    write_internal_outbound_sheet(rows, wb)
    write_zeek_log_sheets(wb, zeek_dfs)
    
    write_legend_sheet(wb)
    
    click.echo("Generating Sankey HTML flows...")
    generate_sankey_html(sankey_data, os.path.join(output_dir, f"{customer_name}_sankey.html"), title="Unfiltered NAVV Purdue Segmentation Flows")
    generate_sankey_html(verified_sankey_data, os.path.join(output_dir, f"{customer_name}_sankey_verified.html"), title="Verified Connections NAVV Purdue Segmentation Flows")
    generate_sankey_html(macro_sankey_data, os.path.join(output_dir, f"{customer_name}_sankey_macro.html"), title="Unfiltered NAVV Macro Purdue Level Flows", link_colors=macro_link_colors)
    generate_sankey_html(verified_macro_sankey_data, os.path.join(output_dir, f"{customer_name}_sankey_macro_verified.html"), title="Verified Connections NAVV Macro Purdue Level Flows", link_colors=macro_link_colors)

    # Generate eleVADR report
    click.echo("Generating eleVADR report...")
    generate_elevadr_report(output_dir, customer_name, rows, inventory, segments)

    click.echo("Auto-adjusting column widths (this may take a minute)...")
    auto_adjust_width(wb["Analysis"])

    click.echo("Calculating capture time statistics...")
    times = (
        perform_zeekcut(fields=["ts"], log_file=os.path.join(zeek_logs, "conn.log"))
        .decode("utf-8")
        .split("\n")[:-1]
    )
    times = [t for t in times if t and not t.startswith("#")]
    forward = sorted(times)
    if not forward:
        error_msg("No connection log data found in conn.log. Please verify that Zeek ran successfully and logs are populated.")
        sys.exit(1)
    try:
        start = float(forward[0])
        end = float(forward[-1])
    except ValueError as e:
        error_msg(f"Failed to parse connection log timestamps: {e}")
        sys.exit(1)
    cap_time = end - start
    timer_data[
        "Length of Capture time"
    ] = "{} day(s) {} hour(s) {} minutes {} seconds".format(
        int(cap_time / 86400),
        int(cap_time % 86400 / 3600),
        int(cap_time % 3600 / 60),
        int(cap_time % 60),
    )
    write_stats_sheet(wb, timer_data)
    write_conn_states_sheet(conn_states, wb)
    write_data_layer_sheet(zeek_df, wb)

    # Reorder sheets to match original layout
    desired_order = [
        "Legend & ReadMe",
        "Analysis",
        inventory_tab_name,
        "Segments",
        "Purdue Violations",
        "External Inbound",
        "Internal Outbound",
        "Internal Hosts",
        "Stats",
        "Conn States",
        "Externals",
        "SNMP",
        "Unknown Internals",
        "Data_layer"
    ]
    if getattr(wb, "vba_archive", None) is not None:
        desired_order.insert(1, "Filters")
    current_sheets = wb.sheetnames
    front_sheets = [s for s in desired_order if s in current_sheets]
    rest = [s for s in current_sheets if s not in front_sheets]
    wb._sheets = [wb[s] for s in front_sheets + rest]

    click.echo("Saving Excel workbook (this can take a few minutes for large datasets)...")
    wb.save(file_name)
    
    # Save external DNS cache for future runs
    with open(ext_dns_path, "w", encoding="utf-8") as f:
        json.dump(ext_dns_cache, f)
    
    # Close geolocator to free resources
    geolocator.close()

    if pcap:
        success_msg(f"Successfully created file: {file_name}")


@click.command("launch")
def launch():
    """Launch the NAVV GUI."""
    port = 5000
    warning_msg("Launching GUI in browser...")
    webbrowser.open(f"http://127.0.0.1:{port}/")
    app.run(port=port)
